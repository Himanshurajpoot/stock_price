"""
Reporting module for generating Excel reports and summaries
"""

import logging
from datetime import datetime
from typing import Dict, Any
import numpy as np
import pandas as pd

logger = logging.getLogger(__name__)


def generate_prediction_report(
    results: Dict[str, Any],
    symbol: str,
    output_path: str = "reports"
) -> str:
    """
    Generate Excel report with predictions and analysis.

    Args:
        results: Results dictionary from run_hybrid_model
        symbol: Stock symbol
        output_path: Directory to save report

    Returns:
        Path to generated report
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils.dataframe import dataframe_to_rows

        # Create Excel writer
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{symbol}_report_{timestamp}.xlsx"

        import os
        os.makedirs(output_path, exist_ok=True)
        filepath = os.path.join(output_path, filename)

        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            # Sheet 1: Summary
            summary_data = {
                "Metric": [
                    "Stock Symbol",
                    "Report Date",
                    "Latest Price",
                    "RMSE",
                    "MAPE",
                    "Direction Accuracy",
                    "XGBoost Weight",
                    "Prophet Weight",
                ],
                "Value": [
                    symbol,
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    f"₹{results['latest_price']:.2f}",
                    results['rmse'],
                    f"{results['mape']}%",
                    f"{results['direction_acc']}%",
                    f"{results['xgb_weight']:.2%}",
                    f"{results['prophet_weight']:.2%}",
                ]
            }
            summary_df = pd.DataFrame(summary_data)
            summary_df.to_excel(writer, sheet_name="Summary", index=False)

            # Sheet 2: Predictions
            pred_df = pd.DataFrame({
                "Date": results["future_dates"],
                "Hybrid Forecast (₹)": np.round(results["hybrid_forecast"], 2),
                "Lower CI (₹)": np.round(results["future_lower_ci"], 2),
                "Upper CI (₹)": np.round(results["future_upper_ci"], 2),
                "XGBoost (₹)": np.round(results["xgb_forecast"], 2),
                "Prophet (₹)": np.round(results["prophet_forecast"], 2),
            })
            pred_df.to_excel(writer, sheet_name="Predictions", index=False)

            # Sheet 3: Technical Indicators (last 30 days)
            indicators_df = pd.DataFrame({
                "RSI": results["rsi"].values,
                "MACD": np.round(results["macd"].values, 4),
                "MACD Signal": np.round(results["macd_signal"].values, 4),
                "BB Upper": np.round(results["bb_upper"].values, 2),
                "BB Middle": np.round(results["bb_middle"].values, 2),
                "BB Lower": np.round(results["bb_lower"].values, 2),
            })
            indicators_df.to_excel(writer, sheet_name="Technical Indicators", index=False)

            # Sheet 4: Backtest Results
            if results.get("backtest"):
                backtest_data = {
                    "Metric": [
                        "Model Return",
                        "Buy & Hold Return",
                        "Sharpe Ratio",
                        "Max Drawdown",
                        "Total Trades",
                        "Final Capital",
                    ],
                    "Value": [
                        f"{results['backtest']['total_return']:.2f}%",
                        f"{results['backtest']['buy_hold_return']:.2f}%",
                        f"{results['backtest']['sharpe_ratio']:.2f}",
                        f"{results['backtest']['max_drawdown']:.2f}%",
                        results['backtest']['trades'],
                        f"₹{results['backtest']['final_capital']:,.2f}",
                    ]
                }
                backtest_df = pd.DataFrame(backtest_data)
                backtest_df.to_excel(writer, sheet_name="Backtest", index=False)

            # Format sheets
            workbook = writer.book

            # Style sheets
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]

                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter

                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass

                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width

                # Format header row
                header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")

                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        logger.info(f"Report generated: {filepath}")
        return filepath

    except Exception as e:
        logger.error(f"Error generating Excel report: {e}")
        return None


def generate_text_summary(results: Dict[str, Any], symbol: str) -> str:
    """
    Generate text summary of results.

    Args:
        results: Results dictionary
        symbol: Stock symbol

    Returns:
        Text summary
    """
    try:
        summary = f"""
╔═══════════════════════════════════════════════════════════╗
║           STOCK PREDICTION ANALYSIS REPORT                ║
╠═══════════════════════════════════════════════════════════╣
║
║ Stock Symbol: {symbol}
║ Report Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
║
║ ─── MARKET DATA ───
║ Latest Price: ₹{results['latest_price']:.2f}
║
║ ─── MODEL PERFORMANCE ───
║ RMSE: {results['rmse']}
║ MAPE: {results['mape']}%
║ Direction Accuracy: {results['direction_acc']}%
║
║ ─── MODEL WEIGHTS ───
║ XGBoost Weight: {results['xgb_weight']:.2%}
║ Prophet Weight: {results['prophet_weight']:.2%}
║ XGBoost RMSE: {results['xgb_rmse']}
║ Prophet RMSE: {results['prophet_rmse']}
"""

        if results.get("backtest"):
            summary += f"""
║ ─── BACKTEST RESULTS ───
║ Model Return: {results['backtest']['total_return']:.2f}%
║ Buy & Hold Return: {results['backtest']['buy_hold_return']:.2f}%
║ Sharpe Ratio: {results['backtest']['sharpe_ratio']:.2f}
║ Max Drawdown: {results['backtest']['max_drawdown']:.2f}%
║ Total Trades: {results['backtest']['trades']}
║ Final Capital: ₹{results['backtest']['final_capital']:,.2f}
"""

        summary += """
╚═══════════════════════════════════════════════════════════╝
"""

        return summary

    except Exception as e:
        logger.error(f"Error generating text summary: {e}")
        return ""


def generate_comparison_report(
    results_list: list,
    output_path: str = "reports"
) -> str:
    """
    Generate comparison report for multiple stocks.

    Args:
        results_list: List of (symbol, results) tuples
        output_path: Directory to save report

    Returns:
        Path to generated report
    """
    try:
        comparison_data = []

        for symbol, results in results_list:
            comparison_data.append({
                "Symbol": symbol,
                "Latest Price": f"₹{results['latest_price']:.2f}",
                "RMSE": results['rmse'],
                "MAPE": f"{results['mape']}%",
                "Direction Acc": f"{results['direction_acc']}%",
                "XGB Weight": f"{results['xgb_weight']:.2%}",
                "Model Return": f"{results['backtest']['total_return']:.2f}%" if results.get('backtest') else "N/A",
            })

        comparison_df = pd.DataFrame(comparison_data)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"comparison_report_{timestamp}.xlsx"

        import os
        os.makedirs(output_path, exist_ok=True)
        filepath = os.path.join(output_path, filename)

        comparison_df.to_excel(filepath, index=False, sheet_name="Comparison")

        logger.info(f"Comparison report generated: {filepath}")
        return filepath

    except Exception as e:
        logger.error(f"Error generating comparison report: {e}")
        return None
